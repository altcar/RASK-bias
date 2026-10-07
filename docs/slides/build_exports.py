"""Compile the two-slide deck to PDF (via headless Chrome) and PPTX (native, editable).

Usage:  .venv/bin/python docs/slides/build_exports.py
Writes: docs/RASK_Bias_Screener.pdf, docs/RASK_Bias_Screener.pptx
"""
import re
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
DOCS = HERE.parent
SLIDES = ["pipeline", "architecture"]
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

# Palette (same as the deck)
INK, PAPER, CARD, LINE = "12182B", "F6F4EF", "FDFCF9", "E2DFD6"
BODY, MUTED = "4A5163", "7A8091"
BLUE, ORANGE = "2563C9", "B9501E"            # on light backgrounds
BLUE_D, ORANGE_D = "6FA3F2", "F08A4B"        # on dark backgrounds
NIGHT_CARD, NIGHT_LINE, NIGHT_BODY = "1C2540", "2E3A5C", "B9C2D6"
SANS, HEAD, MONO = "Calibri", "Calibri", "Courier New"


# ------------------------------------------------------------------------------ PDF
def build_pdf():
    """Wrap the deck's slide HTML in a printable page and print it with headless Chrome."""
    sections = []
    for sid in SLIDES:
        html = (HERE / "project" / "slides" / f"{sid}.html").read_text()
        html = re.sub(r"<aside>.*?</aside>", "", html, flags=re.S)
        html = re.sub(
            r'<x-shape kind="arrow-right" style="background:(#[0-9A-Fa-f]{6});width:(\d+)px;height:(\d+)px"></x-shape>',
            lambda m: (f'<svg width="{m[2]}" height="{m[3]}" viewBox="0 0 10 5" preserveAspectRatio="none" '
                       f'style="flex:none"><polygon points="0,1.5 6,1.5 6,0 10,2.5 6,5 6,3.5 0,3.5" fill="{m[1]}"/></svg>'),
            html)
        sections.append(html)
    page = f"""<!doctype html><html><head><meta charset="utf-8"><title>RASK Bias Screener</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Space+Grotesk:wght@500..700&family=IBM+Plex+Sans:wght@400;600&family=JetBrains+Mono:wght@500&display=swap">
<style>
@page {{ size: 1920px 1080px; margin: 0 }}
* {{ margin: 0; box-sizing: border-box }}
body {{ background: #12182B }}
section {{ width: 1920px; height: 1080px; position: relative; overflow: hidden; page-break-after: always; break-after: page }}
h2, h3 {{ font-weight: 600 }}
p, li {{ line-height: 1.4 }}
ul {{ padding-left: 32px }}
</style></head><body>{''.join(sections)}</body></html>"""
    src = HERE / "print.html"
    src.write_text(page)
    out = DOCS / "RASK_Bias_Screener.pdf"
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    "--virtual-time-budget=15000", f"--print-to-pdf={out}", src.as_uri()],
                   check=True, capture_output=True)
    return out


# ----------------------------------------------------------------------------- PPTX
def rgb(hex_):
    return RGBColor.from_string(hex_)


def box(slide, x, y, w, h, fill, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def arrow(slide, x, y, color, w=0.28, h=0.16):
    return box(slide, x, y - h / 2, w, h, color, shape=MSO_SHAPE.RIGHT_ARROW)


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0.0):
    """paras: list of paragraphs; each is a list of runs (text, size_pt, color, font, bold) + optional space_after."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    for i, para in enumerate(paras):
        runs, after = (para[:-1], para[-1]) if isinstance(para[-1], (int, float)) else (para, 4)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(after)
        for t, size, color, font, bold in runs:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.color.rgb = rgb(color)
            r.font.name = font
            r.font.bold = bold
    return tb


def bullets(slide, x, y, w, h, items, color, size=14):
    tb = text(slide, x, y, w, h, [[(f"•  {it}", size, color, SANS, False), 6] for it in items])
    return tb


def header(slide, eyebrow, title, eyebrow_color, title_color):
    text(slide, 0.89, 0.55, 11.5, 0.35, [[(eyebrow.upper(), 12, eyebrow_color, MONO, True)]])
    text(slide, 0.89, 0.88, 11.6, 0.8, [[(title, 34, title_color, HEAD, True)]])


def footer(slide, label, color):
    text(slide, 0.89, 6.88, 11.6, 0.3, [[(label, 11, color, SANS, False)]])


def slide_pipeline(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(PAPER)
    header(s, "RASK-bias · End to end", "Fair CV screening, from raw data to a live app", BLUE, INK)

    steps = [("01", "Scan", "2,484 Kaggle CVs and 120 US LinkedIn jobs"),
             ("02", "Score", "MiniLM embeds CV and job; cosine is the match"),
             ("03", "Audit", "Same CV, only gender, city or degree swapped"),
             ("04", "Debias", "LEACE erases location and gender signal"),
             ("05", "Ship", "One ONNX model inside a Streamlit app")]
    x, y, w, h, gap = 0.89, 2.0, 1.91, 1.8, 0.5
    for i, (num, title, body) in enumerate(steps):
        cx = x + i * (w + gap)
        box(s, cx, y, w, h, CARD, LINE)
        text(s, cx + 0.16, y + 0.14, w - 0.32, h - 0.28, [
            [(num, 12, BLUE, MONO, True), 2],
            [(title, 20, INK, HEAD, True), 4],
            [(body, 14, BODY, SANS, False)]])
        if i < len(steps) - 1:
            arrow(s, cx + w + 0.11, y + h / 2, "9AA3B5")

    py, ph, pw, pg = 4.15, 2.6, 3.65, 0.3
    panels = [("The model", ORANGE, "all-MiniLM-L6-v2",
               ["6-layer encoder, 384-d vectors", "CVs split into 180-word chunks",
                "PASS if cosine ≥ 0.436 (AUC 0.78)"]),
              ("The debias layer", BLUE, "LEACE concept erasure",
               ["Closed form: x′ = x·A + b", "Fitted on 33 cities, both genders", "Belrose et al., 2023"])]
    for i, (eyebrow, ecol, title, items) in enumerate(panels):
        px = x + i * (pw + pg)
        box(s, px, py, pw, ph, CARD, LINE, radius=0.06)
        text(s, px + 0.25, py + 0.22, pw - 0.5, 0.8, [
            [(eyebrow.upper(), 12, ecol, MONO, True), 4],
            [(title, 20, INK, HEAD, True)]])
        bullets(s, px + 0.25, py + 0.95, pw - 0.5, ph - 1.0, items, BODY)

    px = x + 2 * (pw + pg)
    box(s, px, py, pw, ph, INK, radius=0.06)
    text(s, px + 0.25, py + 0.22, pw - 0.5, 0.35, [[("LOCATION DECISION FLIPS", 12, NIGHT_BODY, MONO, True)]])
    text(s, px + 0.25, py + 0.75, 1.45, 0.7, [[("15.3%", 32, ORANGE_D, HEAD, True)]], anchor=MSO_ANCHOR.MIDDLE)
    arrow(s, px + 1.68, py + 1.1, NIGHT_BODY, w=0.34, h=0.2)
    text(s, px + 2.12, py + 0.75, 1.45, 0.7, [[("10.4%", 32, BLUE_D, HEAD, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, px + 0.25, py + 1.6, pw - 0.5, 0.75,
         [[("Before → after debias. Matching AUC kept: 0.784 → 0.780", 14, NIGHT_BODY, SANS, False)]])

    footer(s, "RASK-bias · UoB Datathon · 1 / 2", MUTED)
    s.notes_slide.notes_text_frame.text = notes("pipeline")


def slide_architecture(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(INK)
    header(s, "Website architecture", "Built once offline, served by one ONNX model", BLUE_D, "EEF0F5")

    lanes = [("Offline", "Python, runs once", ORANGE_D,
              [("data.py", "CVs + US jobs"), ("scan.py", "Gender, location, degree cues"),
               ("audit.py", "Twins, thresholds, metrics"), ("debias.py", "Fit the LEACE eraser"),
               ("export_onnx", "One model, two heads")], 4),
             ("Online", "Streamlit app", BLUE_D,
              [("Inputs", "CV PDF/TXT, job text"), ("Prepare", "Chunk + 6 what-if twins"),
               ("ONNX", "MiniLM + LEACE, 91 MB"), ("Decide", "Cosine vs threshold"),
               ("Show", "Before vs After, bias flag")], 2)]
    x0, w, h, gap = 2.4, 1.66, 1.35, 0.45
    for li, (name, sub, col, boxes, highlight) in enumerate(lanes):
        y = 2.05 + li * 1.6
        text(s, 0.89, y + 0.25, 1.4, 0.8, [[(name, 20, col, HEAD, True), 2], [(sub, 12, NIGHT_BODY, SANS, False)]])
        for i, (title, body) in enumerate(boxes):
            bx = x0 + i * (w + gap)
            box(s, bx, y, w, h, NIGHT_CARD, BLUE_D if i == highlight else NIGHT_LINE)
            text(s, bx + 0.14, y + 0.14, w - 0.28, h - 0.28, [
                [(title, 13, BLUE_D if i == highlight else "EEF0F5", MONO, True), 4],
                [(body, 14, NIGHT_BODY, SANS, False)]])
            if i < len(boxes) - 1:
                arrow(s, bx + w + 0.085, y + h / 2, "5B6888")

    tiles = [("Location flip rate", "15.3%", "10.4%"),
             ("Impact ratio (≥ 0.80 is fair)", "0.85", "0.92"),
             ("Matching AUC (kept)", "0.784", "0.780")]
    ty, tw, th, tg = 5.35, 3.65, 1.3, 0.3
    for i, (label, before, after) in enumerate(tiles):
        tx = 0.89 + i * (tw + tg)
        box(s, tx, ty, tw, th, PAPER, radius=0.08)
        text(s, tx + 0.28, ty + 0.2, tw - 0.56, th - 0.4, [
            [(label, 14, BODY, SANS, False), 4],
            [(before, 28, ORANGE, HEAD, True), ("  →  ", 28, INK, HEAD, True), (after, 28, BLUE, HEAD, True)]])

    footer(s, "Orange = original model (before) · Blue = debiased model (after) · RASK-bias · 2 / 2", "8C96AD")
    s.notes_slide.notes_text_frame.text = notes("architecture")


def notes(sid):
    html = (HERE / "project" / "slides" / f"{sid}.html").read_text()
    m = re.search(r"<aside>(.*?)</aside>", html, flags=re.S)
    return m[1].strip() if m else ""


def build_pptx():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = "RASK Bias Screener"
    slide_pipeline(prs)
    slide_architecture(prs)
    out = DOCS / "RASK_Bias_Screener.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    print(build_pdf())
    print(build_pptx())
